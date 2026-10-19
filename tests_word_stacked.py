"""Word invoices that stack every product inside one table cell."""
import tempfile
from pathlib import Path

from docx import Document

from file_extractor import extract_from_file


def _stacked_invoice(path: Path) -> None:
    doc = Document()
    doc.add_paragraph("CUSTOMS INVOICE")
    doc.add_paragraph("Country of Origin: Great Britain")
    doc.add_paragraph("Our Reference: WCC – ARC 2026 Shipment")
    table = doc.add_table(rows=3, cols=4)
    headers = ("FULL DESCRIPTION OF GOODS", "UNIT VALUE", "TOTAL VALUE", "HS Code")
    for idx, header in enumerate(headers):
        table.rows[0].cells[idx].text = header
    table.rows[1].cells[0].text = "Hand Waving Flags x 1000\nHandbooks x 29\nSales Leaflets printed matter"
    table.rows[1].cells[1].text = "£0.55\n£5.83\n£12.00"
    table.rows[1].cells[2].text = "£550.00\n£169.07\n£12.00"
    table.rows[1].cells[3].text = "4911990000\n4911990000\n4907009000"
    table.rows[2].cells[1].text = "Total Valor / Value"
    table.rows[2].cells[2].text = "£731.07"
    doc.save(path)


def test_stacked_word_cells_become_separate_lines():
    path = Path(tempfile.mkdtemp()) / "arc.docx"
    _stacked_invoice(path)
    _text, items, meta = extract_from_file(open(path, "rb"), path.name, "export")
    goods = [it for it in items if not it.get("_excel_metadata") and "error" not in it]
    assert len(goods) == 3
    flags, books, leaflets = goods
    assert flags["description"] == "Hand Waving Flags x 1000"
    assert flags["quantity"] == 1000
    assert flags["total_value"] == 550
    assert flags["commodity_code"] == "49119900"
    assert flags["country_of_origin"] == "GB"
    assert books["quantity"] == 29
    assert books["total_value"] == 169.07
    assert leaflets["quantity"] == 1
    assert leaflets["total_value"] == 12
    assert abs(sum(it["total_value"] for it in goods) - 731.07) < 0.01
    assert "ARC 2026" in (meta.get("invoice_number") or "")


def test_arc_shipment_file_totals():
    path = Path("/Users/stuartpople/Downloads/ARC Shipment Invoice 2026 01-10-26 (004).docx")
    if not path.exists():
        return
    _text, items, meta = extract_from_file(open(path, "rb"), path.name, "export")
    goods = [it for it in items if not it.get("_excel_metadata") and "error" not in it]
    assert len(goods) == 23
    assert abs(sum(it["total_value"] for it in goods) - 14866.92) < 0.02
    assert all(it["country_of_origin"] == "GB" for it in goods)
    assert abs(float(meta.get("total_invoice_value") or 0) - 14866.92) < 0.02
