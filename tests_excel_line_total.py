"""Spreadsheet invoices with both unit price and line total columns."""
import tempfile
from pathlib import Path

from openpyxl import Workbook

from file_extractor import extract_from_file


def test_wizard_style_uses_line_total_not_unit_price():
    wb = Workbook()
    ws = wb.active
    ws["A1"] = "SHIPPING INVOICE"
    ws["A3"] = "DESCRIPTION"
    ws["B3"] = "COMMODITY CODE"
    ws["C3"] = "COUNTRY OF ORIGIN"
    ws["D3"] = "QTY"
    ws["E3"] = "UNIT WEIGHT KG"
    ws["F3"] = "TOTAL WEIGHT KG"
    ws["G3"] = "UNIT PRICE GBP"
    ws["H3"] = "TOTAL PRICE GBP"
    rows = [
        ("S8 Transfer valve assembly", "8481808590", "UK", 2, 4, 8, 4271, 8542),
        ("S3 Divert valve assembly", "8481808590", "UK", 4, 1.2, 4.8, 1104, 4416),
        ("Optical Tank Sensor Lo M8", None, "UK", 6, 0.026, 0.156, 192, 1152),
    ]
    for i, row in enumerate(rows, start=4):
        for col, val in enumerate(row, start=1):
            ws.cell(i, col, val)
    path = Path(tempfile.mkdtemp()) / "wizard.xlsx"
    wb.save(path)
    _text, items, _meta = extract_from_file(open(path, "rb"), path.name, "export")
    goods = [it for it in items if not it.get("_excel_metadata") and "error" not in it]
    assert len(goods) == 3
    by_desc = {it["description"]: it for it in goods}
    assert by_desc["S8 Transfer valve assembly"]["total_value"] == 8542
    assert by_desc["S8 Transfer valve assembly"]["quantity"] == 2
    assert by_desc["S8 Transfer valve assembly"]["net_weight"] == 8
    assert by_desc["S8 Transfer valve assembly"]["commodity_code"] == "84818085"
    assert by_desc["Optical Tank Sensor Lo M8"]["commodity_code"] == ""
    assert by_desc["Optical Tank Sensor Lo M8"]["total_value"] == 1152
    assert by_desc["Optical Tank Sensor Lo M8"]["needs_review"] is True
    assert abs(sum(it["total_value"] for it in goods) - (8542 + 4416 + 1152)) < 0.01


def test_duplicate_line_and_unit_price_shadow_are_removed():
    """£67,510 was the real £62,680 plus divert valve again (£4,416) and blink unit price (£414)."""
    from file_extractor import _reconcile_to_printed_total

    items = [
        {"description": "S3 Divert valve assembly", "commodity_code": "84818085",
         "quantity": 4, "unit_value": 1104, "total_value": 4416},
        {"description": "S3 Divert valve assembly", "commodity_code": "84818085",
         "quantity": 4, "unit_value": 1104, "total_value": 4416},
        {"description": "8w Blink Panel", "commodity_code": "85371010",
         "quantity": 2, "unit_value": 414, "total_value": 828},
        {"description": "8w Blink Panel", "commodity_code": "85371010",
         "quantity": 2, "unit_value": 414, "total_value": 414},
        {"description": "S3 Scoop Assembly", "commodity_code": "84818081",
         "quantity": 1, "unit_value": 4177.5, "total_value": 4177.5},
    ]
    # Other lines omitted; printed total is the sum of the genuine three lines.
    printed = 4416 + 828 + 4177.5
    kept = [it for it in _reconcile_to_printed_total(items, printed) if not it.get("_excel_metadata")]
    assert abs(sum(it["total_value"] for it in kept) - printed) < 0.05
    assert len(kept) == 3
    assert sorted(it["total_value"] for it in kept) == [828, 4177.5, 4416]
